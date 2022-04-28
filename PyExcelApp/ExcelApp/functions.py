import pandas as pd
import numpy as np
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.workbook import Workbook
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl.chart import BarChart, PieChart, Series, Reference

Filein = input
wb = load_workbook(FileIn)

ws = wb.active

def table (tileset, givename):
    tab = Table(displayName = givename, ref = tileset)
    style = TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False,
    showRowStripes=True, showColumnStripes=True)
    tab.tableStyleInfo = style
    ws.add_table(tab)
    wb.save('new'+Filein+'.xlsx')
